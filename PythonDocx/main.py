"""
从excel文件中读取数据
"""
import openpyxl

def read_from_excel(workbook_path, sheet_name, data_count) -> list:
    """
    :param workbook_path: excel文件路径 data_count: 数据条目数（含表头）
    :return: 数据列表(单条数据以字典形式保存)
    """
    # 打开excel文件
    workbook = openpyxl.load_workbook(workbook_path)
    # 获得指定表单
    sh = workbook[sheet_name]
    # 获得表单数据
    datas = list(sh.rows)
    # 获得表头
    titles = [d.value for d in datas[0]]
    # 存储数据的列表
    infos = []
    # 获得数据并存入字典
    for data in datas[1:data_count]:
        items = [d.value for d in data]
        info = dict(zip(titles, items))
        infos.append(info)
    return infos

"""
写入数据到docx文件
"""
# document = Document(path)
# main_table = document.tables[0]

from docxtpl import DocxTemplate, RichText

def render_by_template(datas, save_dir, template_path, file_name_template, *format_args):
    tpl = DocxTemplate(template_path)
    for data in datas:
        tpl.render(data)
        for arg in format_args:
            file_name_template = file_name_template.format(data[arg])
        tpl.save(save_dir+file_name_template)

"""
主要执行代码
"""
from doc2docx import doc2docx
from xls2xlsx import xls2xlsx

def doc2docx_test():
    docDir = r"D:\Task\学业事务\学工办工作\就业困难精准帮扶"
    #doc2docx_batch(docDir)
    docPath = r"D:\Task\学业事务\学工办工作\就业困难精准帮扶\附件5-就业困难毕业生帮扶情况登记表.doc"
    doc2docx(docPath)

def xls2xlsx_test():
    xlsDir = r"D:\Task\学业事务\学工办工作\就业困难精准帮扶"
    #xls2xlsx_batch(xlsDir)
    xlsPath = r"D:\Task\学业事务\学工办工作\就业困难精准帮扶\dc_sy_yx10500.xls"
    xls2xlsx(xlsPath)

if __name__ == "__main__":
    workbook_path = r"D:\Task\学业事务\学工办工作\就业困难精准帮扶\就业困难.xlsx"
    datas = read_from_excel(workbook_path, "Sheet1", 37)

    workbook_path2 = r"D:\Task\学业事务\学工办工作\就业困难精准帮扶\216876455_0_就业困难毕业生精准帮扶情况登记表_33_33.xlsx"
    datas_2 = read_from_excel(workbook_path, "Sheet1", 37)

    # 读入模板文件
    template_path = r"D:\Task\学业事务\学工办工作\就业困难精准帮扶\附件5-就业困难毕业生帮扶情况登记表.docx"
    save_dir = "D:\\Task\\学业事务\\学工办工作\\就业困难精准帮扶\\登记表\\"
    file_name_template = "{}-{}-{}-就业困难毕业生帮扶情况登记表.docx"
    render_by_template(datas, save_dir, template_path, file_name_template, "学号", "姓名", "专业")
