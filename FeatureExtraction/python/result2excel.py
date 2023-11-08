import os
from pprint import pprint

# 提取数据
filepath = './python/results1k.txt'
export_file = './python/results1k.xlsx'

filepath = os.path.abspath(filepath)
with open(filepath, mode='r', encoding='utf-8') as f:
    lines = f.readlines()

lines = lines[1:-1]
algorithm_line_nums = 39
dataset_line_nums = 7

algo_num = int(len(lines) / algorithm_line_nums)
algorithm_lines = [lines[i*algorithm_line_nums: i*algorithm_line_nums+algorithm_line_nums] for i in range(algo_num)]
algorithms = ["SIFT", "BRISK", "ORB", "AKAZE"]
results = {}

for algo_line in algorithm_lines:
    # get dataset for algo content
    algorithm = algo_line[0][10:-9] # 获得算法名称
    valid_lines = algo_line[1:-3]
    dataset_num = int(len(valid_lines) / dataset_line_nums)
    dataset_results = [valid_lines[i*dataset_line_nums: i*dataset_line_nums+dataset_line_nums] for i in range(dataset_num)]
    
    # save data
    dataset_dict = {}
    for result in dataset_results:
        valid_result = result[0:-2]
        # get data
        dataset = valid_result[0][10:-4]
        total_images = valid_result[1][2:5]
        # TODO: 后续实验都改成英文：了
        avg_time = valid_result[2].split('：')[-1].strip()
        valid_images = valid_result[3].split(':')[-1].strip()
        avg_precision = valid_result[4].split(':')[-1].strip()

        data = {
            'total_iamges': total_images,
            'avg_time': avg_time,
            'valid_images': valid_images,
            'avg_precision': avg_precision
        }

        dataset_dict[dataset] = data
    results[algorithm] = dataset_dict

# pprint(results)

# 制表excel
import openpyxl
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.styles.alignment import Alignment

data = {'Name': ['Alice', 'Bob', 'Charlie', 'David'],
        'Age': [25, 30, 35, 40],
        'City': ['New York', 'Los Angeles', 'Chicago', 'Houston']}


datasets = [key for key in results['SIFT'].keys()]

valid_data, avg_precision_data, avg_time_data = {}, {}, {}

valid_data = {algo: [data[dataset]['valid_images'] for dataset in datasets] for algo, data in results.items()}
avg_time_data = {algo: [data[dataset]['avg_time'] for dataset in datasets] for algo, data in results.items()}
avg_precision_data = {algo: [data[dataset]['avg_precision'] for dataset in datasets] for algo, data in results.items()}

data = [valid_data, avg_time_data, avg_precision_data]

workbook = openpyxl.Workbook()
sheet = workbook.active

# 添加数据到工作表
sheet.append(['有效图样'])
sheet.append(['数据集', *datasets])
for algo, data in valid_data.items():
    sheet.append([algo, *data])

sheet.append([])

sheet.append(['平均运行时间'])
sheet.append(['数据集', *datasets])
for algo, data in avg_time_data.items():
    sheet.append([algo, *data])

sheet.append([])

sheet.append(['平均准确度'])
sheet.append(['数据集', *datasets])
for algo, data in avg_precision_data.items():
    sheet.append([algo, *data])
# 设置样式
columns = ['A', 'B', 'C', 'D', 'E', 'F']
for column in columns:
    sheet.column_dimensions[column].width = 20
# sheet.alignment = Alignment(horizontal='center', vertical='center')
# 合并单元格
merged_cells = ['A1:F1', 'A8:F8', 'A15:F15']
for cells in merged_cells:
    sheet.merge_cells(cells)

# 创建一个数据表格
table = Table(displayName="MyTable", ref="A1:F20")

# 添加数据表格到工作簿
sheet.add_table(table)

# 保存工作簿为Excel文件
workbook.save(export_file)

